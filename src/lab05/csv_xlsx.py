import csv

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """
    Read CSV-file and convert to XLSX-file.

    Input data
    ----------
    csv_path: str
    xlsx_path: str

    Returns
    -------
    None

    Raises
    ------
    ValueError: CSV-file is empty or invalid structure.
    """
    wb = Workbook()
    ws = wb[wb.sheetnames[0]]
    ws.title = "Sheet1"
    wb.active = 0

    with open(csv_path, "r", encoding="utf-8", newline="") as f:
        # Get information about CSV-file
        sample = f.read(1024)
        try:
            dialect = csv.Sniffer().sniff(sample)
        except csv.Error:
            raise ValueError("CSV-file is empty or invalid structure.")

        f.seek(0)
        reader = csv.reader(f, dialect=dialect)

        max_len = []
        for row_idx, row in enumerate(reader, start=1):
            for col_idx, value in enumerate(row, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)

                if col_idx > len(max_len):
                    max_len.append(0)
                max_len[col_idx - 1] = max(max_len[col_idx - 1], len(str(value)))

    # Autowidth (min = 8)
    for i, length in enumerate(max_len, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(length, 8)

    wb.save(xlsx_path)
